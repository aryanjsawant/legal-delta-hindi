from pptx import Presentation

prs = Presentation('Preference-Based Training for Legal Bail Prediction.pptx')
text = []

for idx, slide in enumerate(prs.slides):
    text.append(f"\\n--- Slide {idx + 1} ---\\n")
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            text.append(shape.text)

with open('presentation_extracted.txt', 'w', encoding='utf-8') as f:
    f.write("\\n".join(text))

print("Extraction complete.")
